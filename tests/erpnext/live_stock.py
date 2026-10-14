"""P2–P4（庫存表、全公司補貨、訂單調度）與共用庫存層的整合測試。
在已安裝的正式版 ERPNext 容器內執行，結束時全部復原（不留下測試資料）。

    docker compose cp live_stock.py backend:/tmp/live_stock.py
    docker compose exec backend bash -c "cd sites && ../env/bin/python /tmp/live_stock.py nila"
    （大量測試：再加上款式數，例如 nila 170）
"""

import sys

import frappe

failures = []


def check(label, ok, detail=""):
    print(f"{'✓' if ok else '✗'} {label}{'：' + str(detail) if detail not in ('', None) else ''}", flush=True)
    if not ok:
        failures.append(label)


def raises(fn, *args, **kwargs):
    try:
        fn(*args, **kwargs)
    except Exception as e:  # noqa: BLE001
        frappe.clear_messages()
        return str(e)
    return None


class Env:
    """測試資料：一個款式（兩種甲型 × 兩種尺寸）、一位客戶、各倉庫代碼。"""

    def __init__(self):
        from nila_app import importer, products as p, stock

        self.stock = stock
        whs = stock.warehouses(include_transit=False)
        self.hq = stock.hq_warehouse()
        branches = [w["name"] for w in whs if not w["hq"]]
        self.a, self.b = branches[0], branches[1]
        self.transit = stock.transit_warehouse()
        opts = p.get_style_options()
        self.shapes, self.sizes = opts["shapes"][:2], opts["sizes"][:2]
        res = p.create_style({"name": "測試庫存款", "cost": 100, "web_price": 390, "store_price": 350,
                              "cells": [{"shape": sh, "size": sz} for sh in self.shapes for sz in self.sizes],
                              "targets": {self.shapes[0]: 10, self.shapes[1]: 0}})
        self.style = res["code"]
        self.items = res["variants"]  # [甲型1-尺寸1, 甲型1-尺寸2, 甲型2-尺寸1, 甲型2-尺寸2]
        self.x = self.items[0]
        self.importer = importer
        self.customer = frappe.get_doc({"doctype": "Customer", "customer_name": "測試客戶甲",
                                        "customer_type": "Individual"}).insert().name

    def opening(self, rows):
        """rows：{倉庫: [(料號, 數量)]}"""
        self.importer._opening_stock({wh: [(c, q, 100) for c, q in items] for wh, items in rows.items()})

    def order(self, lines, warehouse=None, submit=True):
        so = frappe.get_doc({
            "doctype": "Sales Order", "customer": self.customer, "transaction_date": frappe.utils.nowdate(),
            "delivery_date": frappe.utils.add_days(frappe.utils.nowdate(), 3), "selling_price_list": "官網售價",
            "set_warehouse": warehouse or self.hq,
            "items": [{"item_code": c, "qty": q, "warehouse": warehouse or self.hq} for c, q in lines],
        }).insert()
        if submit:
            so.submit()
        return so

    def bin(self, item, wh, field="actual_qty"):
        return self.stock.num(frappe.db.get_value("Bin", {"item_code": item, "warehouse": wh}, field))


def test_shared(env):
    s = env.stock
    check("允許負庫存", frappe.db.get_single_value("Stock Settings", "allow_negative_stock") == 1)
    whs = s.warehouses()
    check("倉庫順序：總公司在前、在途在最後", whs[0]["hq"] and whs[-1]["transit"], [w["label"] for w in whs])
    check("英文預設倉庫不在清單", not any(w["label"] in ("Stores", "Goods In Transit") for w in whs))
    check("原因選項有預設值", all(s.reasons(k) for k in s.REASON_KINDS), {k: s.reasons(k) for k in s.REASON_KINDS})
    check("有預設廠商", frappe.db.exists("Supplier", "預設廠商"))

    # 客戶範例：總公司 2、分公司A 1、訂單 5（總公司出貨）→ 全公司預計 −2 → 補 12
    env.opening({env.hq: [(env.x, 2)], env.a: [(env.x, 1)]})
    so = env.order([(env.x, 5)])
    b = s.bins([env.x])[env.x]
    total = sum(v["projected_qty"] for v in b.values())
    check("總公司預計庫存 −3", b[env.hq]["projected_qty"] == -3, b[env.hq])
    check("全公司預計庫存 −2", total == -2, total)
    from nila_app.stockrules import replenish_qty
    check("建議補貨 12", replenish_qty(10, total, "上架") == 12)

    lines, res, tr = s.allocation([env.x])
    r = res[lines[0]["key"]]
    check("分配：現貨 2、調撥 1、要叫貨 2", (r["stock"], r["transfer"], r["short"]) == (2, 1, 2), r)
    check("調撥建議：分公司A → 總公司 1 件", [(t["from"], t["to"], t["qty"]) for t in tr] == [(env.a, env.hq, 1)], tr)

    # 在途調撥：分公司A 寄出 1 件 → 在途 → 總公司入庫
    err = raises(s.create_transfer, env.a, env.a, [(env.x, 1)])
    check("調出與調入相同時拒絕", err and "不能相同" in err, err)
    err = raises(s.create_transfer, env.a, env.transit, [(env.x, 1)])
    check("不能直接調到在途倉", err and "使用中" in err, err)
    out = s.create_transfer(env.a, env.hq, [(env.x, 1)], sales_order=so.name)
    check("寄出後分公司A 為 0、在途 1", (env.bin(env.x, env.a), env.bin(env.x, env.transit)) == (0, 1))
    check("在途數量記在總公司名下", s.transit_inbound([env.x]) == {(env.x, env.hq): 1}, s.transit_inbound([env.x]))
    lines, res, tr = s.allocation([env.x])
    r = res[lines[0]["key"]]
    check("分配：在途 1 抵掉調撥建議", (r["stock"], r["transit"], r["transfer"], r["short"]) == (2, 1, 0, 2) and not tr, r)
    total = sum(v["projected_qty"] for v in s.bins([env.x])[env.x].values())
    check("調撥不改變全公司合計（仍補 12）", replenish_qty(10, total, "上架") == 12, total)

    inn = s.receive_transfer(out)
    check("到貨入庫：總公司 3、在途 0", (env.bin(env.x, env.hq), env.bin(env.x, env.transit)) == (3, 0))
    check("入庫單記錄相關訂單", frappe.db.get_value("Stock Entry", inn, "nila_sales_order") == so.name)
    check("已入庫的調撥單不在在途清單", s.transit_inbound([env.x]) == {})
    err = raises(s.receive_transfer, out)
    check("不能重複入庫", err and "已經全部入庫" in err, err)

    # 調整原因紀錄
    err = raises(s.log_reason, "補貨調整", env.x, 12, 20, "")
    check("沒有原因時拒絕", err and "請選擇原因" in err, err)
    s.log_reason("負庫存", env.x, None, None, "訂單超賣", note="測試")
    check("最新原因", s.latest_reasons("負庫存", [env.x])[env.x].reason == "訂單超賣")
    return so


def find(blocks, code):
    return next((b for b in blocks if b["code"] == code), None)


def row(block, label):
    return next((r for r in block["rows"] if r["label"] == label), None)


def test_matrix(env):
    from nila_app import matrix as m
    from nila_app.stockrules import LOW, NEGATIVE, OUT

    x, y, z = env.x, env.items[1], env.items[2]
    # 分公司B 賣出 1 件 y（帳上沒有）→ 負庫存
    se = frappe.get_doc({"doctype": "Stock Entry", "stock_entry_type": "Material Issue", "purpose": "Material Issue",
                         "items": [{"item_code": y, "qty": 1, "s_warehouse": env.b}]}).insert()
    se.submit()
    check("允許負庫存：分公司B 變 −1", env.bin(y, env.b) == -1)

    blocks, whs = m.build({"q": "測試庫存款"})
    b = find(blocks, env.style)
    check("找得到款式區塊", b is not None)
    labels = [r["label"] for r in b["rows"]]
    check("列：總公司、各分公司、全公司合計、建議補貨（沒有在途就不顯示在途）",
          labels[0] == "總公司" and labels[-2:] == ["全公司合計", "建議補貨"] and "在途" not in labels, labels)
    check("欄：兩種甲型，每種顯示全部尺寸，沒有的規格留白",
          [c["shape"] for c in b["columns"]] == env.shapes
          and [z_["item"] for z_ in b["columns"][0]["sizes"]][:2] == env.items[:2]
          and all(z_["item"] is None for z_ in b["columns"][0]["sizes"][2:]), b["columns"][0])
    total = row(b, "全公司合計")["cells"]
    check("實際庫存：x 合計 3、低於目標 10＝低庫存", total[x] == {"q": 3, "s": LOW}, total[x])
    check("沒有庫存、沒設目標：缺貨", total[z]["s"] == OUT, total[z])
    check("y 合計 −1＝負庫存", total[y] == {"q": -1, "s": NEGATIVE}, total[y])
    check("分公司B 的 y 標負庫存", row(b, "分公司B")["cells"][y]["s"] == NEGATIVE)
    check("分公司的 0 不標缺貨", row(b, "分公司A")["cells"][x]["s"] is None)
    check("建議補貨：x 12、y 11（目標 10 − (−1)）、甲型2 不補",
          (row(b, "建議補貨")["cells"][x]["q"], row(b, "建議補貨")["cells"][y]["q"], row(b, "建議補貨")["cells"][z]["q"]) == (12, 11, ""))
    check("小計", row(b, "總公司")["sum"] == 3 and row(b, "全公司合計")["sum"] == 2, [r["sum"] for r in b["rows"]])

    b = find(m.build({"q": env.style, "mode": "projected"})[0], env.style)
    check("預計庫存：x 合計 −2＝負庫存", row(b, "全公司合計")["cells"][x] == {"q": -2, "s": NEGATIVE})

    b = find(m.build({"q": env.style, "shape": env.shapes[0]})[0], env.style)
    check("篩選甲型：只剩一種甲型", [c["shape"] for c in b["columns"]] == env.shapes[:1])
    b = find(m.build({"q": env.style, "warehouse": env.a})[0], env.style)
    check("篩選公司：只顯示該公司與合計", [r["label"] for r in b["rows"]][:2] == ["分公司A", "全公司合計"], [r["label"] for r in b["rows"]])
    check("只看需補貨：包含此款", find(m.build({"q": env.style, "only_need": 1})[0], env.style) is not None)
    check("庫存狀態篩選：負庫存", find(m.build({"q": env.style, "state": "負庫存"})[0], env.style) is not None)
    check("商品狀態篩選：停產不含此款", find(m.build({"q": env.style, "status": "停產"})[0], env.style) is None)
    check("搜尋不到時為空", m.build({"q": "不存在的款式zzz"})[0] == [])

    r = m.get_matrix({"q": env.style}, 0, 1)
    check("分頁與摘要", r["total"] >= 1 and len(r["styles"]) == 1 and r["summary"]["need_qty"] >= 23, r["summary"])
    m.export_matrix({"q": env.style})
    from io import BytesIO
    import openpyxl
    ws = openpyxl.load_workbook(BytesIO(frappe.response["filecontent"])).active
    vals = list(ws.values)
    check("匯出 Excel：表頭與資料列", vals[0][:4] == ("款式編號", "款式名稱", "商品狀態", "公司")
          and any(v[0] == env.style and v[3] == "全公司合計" for v in vals[1:]), vals[:2])
    frappe.response.pop("filecontent", None)


def status(so):
    return frappe.db.get_value("Sales Order", so if isinstance(so, str) else so.name, "nila_dispatch_status")


def sugg(env, item):
    from nila_app import replenish

    rows, moves = replenish.suggestions(q=env.style)
    return next((r for r in rows if r["item"] == item), None), moves


def test_replenish(env, so1):
    from nila_app import products as p
    from nila_app import replenish
    from erpnext.buying.doctype.purchase_order.purchase_order import make_purchase_receipt

    x, y = env.x, env.items[1]
    r, moves = sugg(env, x)
    check("補貨建議：x 補 12（目標 10 − 全公司預計 −2）", r and r["suggest"] == 12, r)
    ry, _ = sugg(env, y)
    check("補貨建議：y 補 11 且標示負庫存", ry and ry["suggest"] == 11 and ry["actual"] == -1, ry)
    check("沒設目標、沒有負庫存的規格不列出", sugg(env, env.items[2])[0] is None)
    check("沒有其他公司有貨時沒有調撥建議", moves == [], moves)

    env.opening({env.a: [(x, 1)]})  # 分公司A 又有 1 件 → 建議調回總公司
    r, moves = sugg(env, x)
    check("其他公司有貨：建議分公司A → 總公司 1 件（給訂單 1）",
          [(m["from"], m["to"], m["qty"], m["orders"]) for m in moves] == [(env.a, env.hq, 1, [so1.name])], moves)
    check("調撥不影響補貨建議（全公司合計不變，補 11）", r["suggest"] == 11, r)

    err = raises(replenish.create_material_request, [{"item": x, "qty": 15, "suggest": 11}])
    check("叫貨數和建議不同又沒選原因：拒絕", err and "請選擇原因" in err, err)
    mr = replenish.create_material_request([{"item": x, "qty": 15, "suggest": 11, "reason": "預計熱賣", "note": "週年慶"},
                                           {"item": y, "qty": 11, "suggest": 11}])
    check("建立補貨單（待叫貨）", frappe.db.get_value("Material Request", mr, ["docstatus", "status"]) == (1, "Pending"))
    log = frappe.get_all("Nila Reason Log", filters={"reference_name": mr}, fields=["item_code", "before_qty", "after_qty", "reason", "kind"])
    check("只有改過數量的規格留下調整紀錄", [(l.item_code, l.before_qty, l.after_qty, l.reason, l.kind) for l in log]
          == [(x, 11, 15, "預計熱賣", "補貨調整")], log)
    r, _ = sugg(env, x)
    check("建立補貨單後建議數變 0（不會重複叫貨）", r is None or r["suggest"] == 0, r)

    stopped = p.create_style({"name": "測試停產款", "cost": 1, "web_price": 1, "store_price": 1, "status": "停產",
                              "cells": [{"shape": env.shapes[0], "size": env.sizes[0]}], "targets": {env.shapes[0]: 5}})
    err = raises(replenish.create_material_request, [{"item": stopped["variants"][0], "qty": 1, "suggest": 1}])
    check("停產的商品不能建立補貨單", err and "已停產" in err, err)

    po = replenish.mark_ordered(mr)
    check("已叫貨：建立採購單、補貨單變已叫貨",
          frappe.db.get_value("Purchase Order", po, "docstatus") == 1 and frappe.db.get_value("Material Request", mr, "status") == "Ordered")
    check("採購單單價＝成本", frappe.db.get_value("Purchase Order Item", {"parent": po, "item_code": x}, "rate") == 100)
    check("已叫貨數量計入預計庫存", env.bin(x, env.hq, "ordered_qty") == 15 and env.bin(x, env.hq, "indented_qty") == 0)
    err = raises(replenish.mark_ordered, mr)
    check("不能重複叫貨", err and "不是待叫貨" in err, err)

    r, moves = sugg(env, x)
    names = replenish.create_transfers([dict(moves[0], qty=1, suggest=1)])
    check("建立調撥單：分公司A → 在途，記錄訂單", frappe.db.get_value("Stock Entry", names[0], ["nila_destination", "nila_sales_order"])
          == (env.hq, so1.name) and env.bin(x, env.a) == 0)
    check("調撥後沒有新的調撥建議", sugg(env, x)[1] == [])
    err = raises(replenish.create_transfers, [{"item": x, "from": env.hq, "to": env.a, "qty": 2, "suggest": 1}])
    check("調撥數和建議不同又沒選原因：拒絕", err and "請選擇原因" in err, err)

    replenish.set_negative_reason(y, "門市先賣未入帳", "週六市集")
    ry, _ = sugg(env, y)
    check("負庫存原因顯示在建議清單", ry and ry["negative_reason"] == "門市先賣未入帳：週六市集", ry)

    check("訂單 1 狀態：差的由已叫貨補上＝已叫貨", status(so1) == "已叫貨", status(so1))
    pr = make_purchase_receipt(po)
    pr.insert()
    pr.submit()
    check("到貨入庫：總公司 +15、補貨單已到貨", env.bin(x, env.hq) == 18 and frappe.db.get_value("Material Request", mr, "status") == "Received",
          (env.bin(x, env.hq), frappe.db.get_value("Material Request", mr, "status")))
    check("到貨後訂單 1 變可出貨", status(so1) == "可出貨", status(so1))

    before = frappe.db.count("Notification Log", {"link": "/app/nila-replenish"})
    replenish.daily()
    check("每日：有需要補貨時通知管理者", frappe.db.count("Notification Log", {"link": "/app/nila-replenish"}) > before)
    return mr


def test_dispatch(env, so1):
    from nila_app import dispatch, replenish
    from erpnext.selling.doctype.sales_order.sales_order import make_delivery_note

    x = env.x
    hq_free = env.bin(x, env.hq) - 5  # 總公司扣掉訂單 1 剩下的
    err = raises(env.order, [(x, 1)], warehouse=env.transit)
    check("出貨倉不能是在途倉", err and "出貨倉必須是" in err, err)

    draft = {"docstatus": 0, "items": [{"idx": 1, "item_code": x, "qty": 2, "warehouse": env.a, "conversion_factor": 1}]}
    pnl = dispatch.order_panel(draft)
    res = pnl["lines"][0]["result"]
    check("訂單畫面（未送出、分公司A 出貨）：從總公司調 2", (res["stock"], res["transfer"], res["sources"]) == (0, 2, [(env.hq, 2)]), res)
    check("訂單畫面：狀態待調撥、列出各倉庫存", pnl["status"] == "待調撥" and pnl["lines"][0]["stock"][env.hq] == env.bin(x, env.hq), pnl["status"])

    so2 = env.order([(x, 2)], warehouse=env.a)
    check("分公司出貨的訂單：待調撥", status(so2) == "待調撥", status(so2))
    check("先來的訂單 1 不受影響", status(so1) == "可出貨")
    pnl = dispatch.order_panel(so2.as_dict())
    check("已送出訂單的調撥建議：總公司 → 分公司A 2 件", [(t["from"], t["to"], t["qty"], t["orders"]) for t in pnl["transfers"]]
          == [(env.hq, env.a, 2, [so2.name])], pnl["transfers"])
    se = replenish.create_transfers([dict(t, qty=t["qty"], suggest=t["qty"]) for t in pnl["transfers"]])[0]
    check("調撥在途：訂單仍待調撥", status(so2) == "待調撥")
    check("訂單畫面列出在途調撥單", [m["name"] for m in dispatch.order_panel(so2.as_dict())["moving"]] == [se])
    from nila_app import stock
    stock.receive_transfer(se)
    check("調撥到貨後訂單變可出貨", status(so2) == "可出貨", status(so2))

    dn = make_delivery_note(so2.name)
    dn.insert()
    dn.submit()
    check("出貨後：已出貨、分公司A 庫存扣掉", status(so2) == "已出貨" and env.bin(x, env.a) == 0, (status(so2), env.bin(x, env.a)))
    check("出貨後預留歸零", env.bin(x, env.a, "reserved_qty") == 0)

    so3 = env.order([(x, hq_free + 1)])
    check("總公司不夠、其他公司也沒有：待叫貨", status(so3) == "待叫貨", status(so3))
    so4 = env.order([(x, 1)])
    check("排在後面的訂單也缺：待叫貨", status(so4) == "待叫貨", status(so4))
    so3.reload()
    so3.cancel()
    check("取消訂單 3：狀態清空、後面的訂單 4 變可出貨", (status(so3), status(so4)) == ("", "可出貨"), (status(so3), status(so4)))

    from erpnext.controllers.accounts_controller import update_child_qty_rate
    so4.reload()
    import json
    update_child_qty_rate("Sales Order", json.dumps([{"docname": so4.items[0].name, "item_code": x, "qty": hq_free + 50,
                                                      "rate": so4.items[0].rate, "conversion_factor": 1}]), so4.name)
    check("修改訂單數量（更新明細）後重算：待叫貨", status(so4) == "待叫貨", status(so4))

    so5 = env.order([(x, 1)], warehouse=env.b)
    t = dispatch.order_panel(so5.as_dict())["transfers"]
    check("分公司B 的訂單要從其他公司調：沒有可調的庫存時待叫貨", status(so5) == "待叫貨" and t == [], (status(so5), t))
    se = stock.create_transfer(env.hq, env.b, [(x, 1)], sales_order=so5.name)
    frappe.clear_messages()
    so5.reload()
    so5.cancel()
    msgs = " ".join(str(m) for m in frappe.message_log)
    check("取消有在途調撥的訂單：提醒調撥單仍在途", se in msgs and "仍在途" in msgs, msgs[:200])
    frappe.clear_messages()

    # 門市銷售（庫存變動）也會重算：總公司領出到剩不夠訂單 1
    so1.reload()
    issue = frappe.get_doc({"doctype": "Stock Entry", "stock_entry_type": "Material Issue", "purpose": "Material Issue",
                            "items": [{"item_code": x, "qty": env.bin(x, env.hq), "s_warehouse": env.hq}]}).insert()
    issue.submit()
    check("庫存被領走後訂單 1 重算為待叫貨", status(so1) == "待叫貨", status(so1))
    frappe.db.set_value("Sales Order", so1.name, "nila_dispatch_status", "可出貨")
    dispatch.refresh_all()
    check("每日全部重算會修正狀態", status(so1) == "待叫貨")


def run_scale(n_styles, n_orders=300):
    """大量測試：n 款 × 7 種甲型 × 全部尺寸，每個規格每個倉庫都有庫存，再加 n_orders 張未出貨訂單（每張 3 行）。
    資料直接寫進資料表（只為測速度，結束時復原），量測庫存表、補貨建議、訂單狀態重算、訂單畫面的時間。"""
    import random
    import time

    from nila_app import dispatch, matrix, replenish, stock

    attr = frappe.get_doc("Item Attribute", "甲型")
    for name, ab in [("測試甲型", "XA"), ("測試甲型二", "XB"), ("測試甲型三", "XC"), ("測試甲型四", "XD")][: max(0, 7 - len(attr.item_attribute_values))]:
        attr.append("item_attribute_values", {"attribute_value": name, "abbr": ab})
    attr.save()
    shapes, sizes = list(stock.shapes()), stock.sizes()
    whs = [w["name"] for w in stock.warehouses()]
    now = frappe.utils.now()
    base = {"creation": now, "modified": now, "owner": "Administrator", "modified_by": "Administrator", "docstatus": 0}
    items, bins, variants = [], [], []
    for i in range(n_styles):
        code = f"S{i:04d}"
        items.append(dict(base, name=code, item_code=code, item_name=f"速度測試{i:03d}", item_group="穿戴甲", stock_uom="件",
                          has_variants=1, is_stock_item=1, nila_status="上架"))
        for sh in shapes:
            for sz in sizes:
                v = f"{code}-{shapes.index(sh)}-{sz}"
                variants.append(v)
                items.append(dict(base, name=v, item_code=v, item_name=f"速度測試{i:03d}-{sh}-{sz}", item_group="穿戴甲",
                                  stock_uom="件", variant_of=code, is_stock_item=1, nila_shape=sh, nila_size=sz,
                                  nila_status="上架", nila_reorder_target=10))
                for w in whs:
                    q = random.randint(-1, 5)
                    bins.append(dict(base, name=frappe.generate_hash(length=12), item_code=v, warehouse=w, actual_qty=q,
                                     projected_qty=q, reserved_qty=0, ordered_qty=0, indented_qty=0, stock_uom="件"))
    keys = sorted({k for it in items for k in it})
    defaults = {"has_variants": 0, "nila_reorder_target": 0}
    frappe.db.bulk_insert("Item", keys, [[it.get(k, defaults.get(k)) for k in keys] for it in items], chunk_size=2000)
    bkeys = list(bins[0].keys())
    frappe.db.bulk_insert("Bin", bkeys, [[b[k] for k in bkeys] for b in bins], chunk_size=5000)
    customer = frappe.get_doc({"doctype": "Customer", "customer_name": "速度測試客戶", "customer_type": "Individual"}).insert().name
    company = frappe.db.get_value("Company", {}, "name")
    sos, sois = [], []
    for n in range(n_orders):
        so = f"SPEED-SO-{n:05d}"
        sos.append(dict(base, name=so, docstatus=1, status="To Deliver and Bill", customer=customer, company=company,
                        transaction_date=frappe.utils.nowdate(), per_delivered=0, currency="TWD", conversion_rate=1,
                        selling_price_list="官網售價", price_list_currency="TWD", plc_conversion_rate=1, order_type="Sales"))
        for idx in range(1, 4):
            sois.append(dict(base, name=f"{so}-{idx}", parent=so, parenttype="Sales Order", parentfield="items", idx=idx,
                             docstatus=1, item_code=random.choice(variants), warehouse=random.choice(whs[:-1]), qty=random.randint(1, 3),
                             delivered_qty=0, conversion_factor=1, uom="件", stock_uom="件"))
    k = list(sos[0].keys())
    frappe.db.bulk_insert("Sales Order", k, [[s[x] for x in k] for s in sos])
    k = list(sois[0].keys())
    frappe.db.bulk_insert("Sales Order Item", k, [[s[x] for x in k] for s in sois])
    print(f"  （資料：{n_styles} 款、{len(variants)} 個規格、{len(bins)} 筆庫存、{n_orders} 張訂單 {len(sois)} 行）")

    def timed(label, fn, limit):
        t = time.time()
        out = fn()
        used = time.time() - t
        check(f"大量：{label}（{used:.2f} 秒，目標 {limit} 秒內）", used < limit)
        return out

    r = timed("庫存表第一頁", lambda: matrix.get_matrix({}, 0, 20), 3)
    check("大量：庫存表款式數", r["total"] >= n_styles and len(r["styles"]) == 20, r["total"])
    timed("庫存表只看需補貨", lambda: matrix.get_matrix({"only_need": 1, "mode": "projected"}, 0, 20), 3)
    timed("庫存表搜尋一款", lambda: matrix.get_matrix({"q": "速度測試150"}, 0, 20), 3)
    timed("庫存表匯出 Excel（全部）", lambda: matrix.export_matrix({}), 15)
    frappe.response.pop("filecontent", None)
    rows, moves = timed("補貨建議（全部，含調撥建議）", lambda: replenish.suggestions(), 5)
    check("大量：有補貨建議與調撥建議", len(rows) > 100 and len(moves) > 0, (len(rows), len(moves)))
    timed("每日重算全部訂單出貨狀態", dispatch.refresh_all, 10)
    check("大量：訂單都有出貨狀態", frappe.db.count("Sales Order", {"name": ["like", "SPEED-SO-%"], "nila_dispatch_status": ["is", "set"]}) == n_orders)
    doc = {"docstatus": 0, "items": [{"idx": i, "item_code": variants[i * 7], "qty": 2, "warehouse": whs[0]} for i in range(1, 11)]}
    timed("訂單畫面（10 行，未送出）", lambda: dispatch.order_panel(doc), 2)
    timed("庫存變動後重算相關訂單（30 個規格）", lambda: dispatch.refresh_items(variants[:30]), 2)


def run():
    env = Env()
    so1 = test_shared(env)
    test_matrix(env)
    test_replenish(env, so1)
    test_dispatch(env, so1)


def main():
    site = sys.argv[1] if len(sys.argv) > 1 else "nila"
    frappe.init(site=site, sites_path=".")
    frappe.connect()
    frappe.set_user("Administrator")
    try:
        if len(sys.argv) > 2:
            run_scale(int(sys.argv[2]))
        else:
            run()
    except Exception as e:  # noqa: BLE001
        import traceback
        traceback.print_exc()
        check("執行過程沒有錯誤", False, e)
    finally:
        frappe.db.rollback()
        frappe.destroy()
    print(f"\n{'全部通過' if not failures else '未通過：' + '、'.join(failures)}")
    return 0 if not failures else 1


if __name__ == "__main__":
    sys.exit(main())
